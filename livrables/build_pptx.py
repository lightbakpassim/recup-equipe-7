from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pathlib import Path

OUT = Path(__file__).resolve().parent / "Recup-presentation.pptx"
OUT_ALT = Path(__file__).resolve().parent / "Recup-presentation-EQUIPE_7.pptx"

INK = RGBColor(0x17, 0x13, 0x10)
PAPER = RGBColor(0xF1, 0xEA, 0xDC)
BRICK = RGBColor(0xC4, 0x3C, 0x28)
SAND = RGBColor(0xE5, 0xD6, 0xBF)
CARD = RGBColor(0xFF, 0xFA, 0xF2)
MUTED = RGBColor(0x6B, 0x62, 0x58)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def bg(slide, color=PAPER):
    fill = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    fill.line.fill.background()
    fill.fill.solid()
    fill.fill.fore_color.rgb = color
    spTree = slide.shapes._spTree
    sp = fill._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def box(slide, l, t, w, h, color, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, l, t, w, h)
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    return shape


def text(slide, l, t, w, h, content, size=18, bold=False, color=INK, font="Calibri", align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = content
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return tb


def kicker(slide, label):
    text(slide, Inches(0.8), Inches(0.45), Inches(11), Inches(0.4), label.upper(), 13, True, BRICK)


MEMBERS = [
    "BAKPASSIM Pouwedeo Light",
    "ABOLO-SEWOVI Ami Raphaëlla",
    "KOUYAKOUTOULI Godwin Marc",
]

# 1
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
kicker(s, "ESIG Tech Arena 2026  ·  Mini-challenge  ·  EQUIPE_7")
text(s, Inches(0.8), Inches(1.35), Inches(12), Inches(1.4), "Récup’", 80, True, INK, "Georgia")
text(
    s,
    Inches(0.8),
    Inches(2.9),
    Inches(11),
    Inches(1.1),
    "Le tableau du campus pour retrouver ce que tu as perdu,\net emprunter ce qu’il te manque.",
    24,
    False,
    INK,
)
pill = box(s, Inches(0.8), Inches(4.25), Inches(2.2), Inches(0.42), BRICK, True)
tf = pill.text_frame
tf.paragraphs[0].alignment = PP_ALIGN.CENTER
run = tf.paragraphs[0].add_run()
run.text = "EQUIPE_7"
run.font.size = Pt(14)
run.font.bold = True
run.font.color.rgb = PAPER
run.font.name = "Calibri"
for i, name in enumerate(MEMBERS):
    text(s, Inches(0.85), Inches(4.85 + i * 0.38), Inches(11), Inches(0.38), name, 18, False, INK)

# 2
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
kicker(s, "Le problème")
text(s, Inches(0.8), Inches(1.3), Inches(11.5), Inches(2), "Badge perdu. WhatsApp saturé.\nObjet jamais rendu.", 40, True, INK, "Georgia")
text(
    s,
    Inches(0.8),
    Inches(3.8),
    Inches(11),
    Inches(2.2),
    "Aujourd’hui l’info est éclatée : groupes, bouche-à-oreille, accueil.\nOn perd du temps, de l’argent, et beaucoup d’objets restent dans un sac.",
    22,
    False,
    MUTED,
)

# 3
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
kicker(s, "Contexte & utilisateurs")
text(s, Inches(0.8), Inches(1.15), Inches(12), Inches(1.3), "Trois étudiants, le même campus,\nzéro canal commun.", 36, True, INK, "Georgia")
cards = [
    ("Lina", "A perdu son badge entre l’amphi et la cafet."),
    ("Karim", "L’a trouvé. Ne sait pas à qui le rendre."),
    ("Yassine", "Peut prêter sa calculatrice avant l’exam."),
]
for i, (title, body) in enumerate(cards):
    card = box(s, Inches(0.8 + i * 4.05), Inches(3.4), Inches(3.8), Inches(2.4), CARD, True)
    text(s, Inches(1.05 + i * 4.05), Inches(3.65), Inches(3.3), Inches(0.6), title, 24, True, INK)
    text(s, Inches(1.05 + i * 4.05), Inches(4.35), Inches(3.3), Inches(1.2), body, 16, False, MUTED)

# 4
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
kicker(s, "La solution")
text(s, Inches(0.8), Inches(1.3), Inches(12), Inches(2), "Un tableau unique.\nTrois gestes. Trente secondes.", 40, True, INK, "Georgia")
for i, label in enumerate(["Trouvé", "Perdu", "Prêt"]):
    pill = box(s, Inches(0.8 + i * 2.1), Inches(4.0), Inches(1.9), Inches(0.5), INK, True)
    tf = pill.text_frame
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run = tf.paragraphs[0].add_run()
    run.text = label
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = PAPER
text(s, Inches(0.8), Inches(4.9), Inches(11), Inches(1), "On publie, on se donne RDV sur le campus, on clôture. Le fil reste propre.", 20, False, MUTED)

# 5
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
kicker(s, "Fonctionnalités")
feats = [
    ("1. J’ai trouvé", "Photo, catégorie, lieu."),
    ("2. J’ai perdu", "Recherche visible, pas un message noyé."),
    ("3. Je prête", "Calculatrice, chargeur, HDMI."),
    ("4. C’est à moi", "RDV BU, cafet, hall."),
    ("5. Récupéré", "L’annonce se ferme."),
]
for i, (title, body) in enumerate(feats):
    col = i if i < 3 else i - 3
    row = 0 if i < 3 else 1
    x = Inches(0.8 + col * 4.05)
    y = Inches(1.5 + row * 2.5)
    box(s, x, y, Inches(3.8), Inches(2.2), CARD, True)
    text(s, x + Inches(0.25), y + Inches(0.3), Inches(3.3), Inches(0.6), title, 22, True, INK)
    text(s, x + Inches(0.25), y + Inches(1.0), Inches(3.3), Inches(0.8), body, 16, False, MUTED)

# 6
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
kicker(s, "Prototype")
text(s, Inches(0.55), Inches(0.85), Inches(12), Inches(0.7), "Du badge trouvé au RDV.", 28, True, INK, "Georgia")
caps = Path(__file__).resolve().parent / "captures"
shots = ["01-accueil.png", "04-fiche.png", "06-recupere.png"]
for i, name in enumerate(shots):
    img = caps / name
    if img.exists():
        s.shapes.add_picture(str(img), Inches(0.45 + i * 4.25), Inches(1.7), Inches(4.05), Inches(4.7))
text(s, Inches(0.55), Inches(6.42), Inches(10), Inches(0.4), "On vous montre le parcours en live.", 16, False, MUTED)

# 7
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
kicker(s, "Valeur & perspectives")
text(s, Inches(0.8), Inches(1.3), Inches(12), Inches(2), "Moins de rachat. Moins de stress.\nPlus d’entraide.", 40, True, INK, "Georgia")
text(
    s,
    Inches(0.8),
    Inches(3.8),
    Inches(11.5),
    Inches(2),
    "Si on est retenus : alertes de similarité, carte des lieux, dépôt accueil.\nLe cœur du produit est déjà là.",
    22,
    False,
    MUTED,
)

# 8
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
kicker(s, "Conclusion")
text(s, Inches(0.8), Inches(1.8), Inches(12), Inches(1.5), "Récup’.", 80, True, INK, "Georgia")
text(
    s,
    Inches(0.8),
    Inches(3.6),
    Inches(11),
    Inches(2.2),
    "Un problème que chacun a déjà vécu.\nUne solution qu’on explique en une phrase.\nUn proto qu’on vient de parcourir.\n\nMerci — on répond à vos questions.",
    22,
    False,
    INK,
)

total = len(prs.slides)
for i, slide in enumerate(prs.slides, start=1):
    text(slide, Inches(0.8), Inches(7.08), Inches(7), Inches(0.28), "EQUIPE_7  ·  Récup’", 11, False, MUTED)
    text(
        slide,
        Inches(10.3),
        Inches(7.08),
        Inches(2.2),
        Inches(0.28),
        f"{i} / {total}",
        12,
        True,
        MUTED,
        align=PP_ALIGN.RIGHT,
    )

try:
    prs.save(OUT)
    print(OUT)
except PermissionError:
    prs.save(OUT_ALT)
    print(f"LOCKED {OUT}")
    print(OUT_ALT)
