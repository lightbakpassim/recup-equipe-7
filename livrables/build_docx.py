from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE / "document-explicatif.docx"
CAPS = HERE / "captures"

INK = RGBColor(0x17, 0x13, 0x10)
BRICK = RGBColor(0xC4, 0x3C, 0x28)
MUTED = RGBColor(0x5E, 0x56, 0x4D)

SHOTS = [
    ("01-accueil.png", "Écran 1 — Accueil. Lina comprend le produit en une seconde : trois gestes, et les objets près d’elle."),
    ("02-annonces.png", "Écran 2 — Fil Trouvé. Des cartes visuelles : photo, lieu, heure. On scanne, on ne lit pas un roman."),
    ("03-publier.png", "Écran 3 — Karim publie le badge trouvé à la cafet. Quatre champs, trente secondes."),
    ("04-fiche.png", "Écran 4 — Fiche objet. Lina reconnaît son badge et clique « C’est à moi »."),
    ("05-rdv.png", "Écran 5 — RDV campus. Lieu + heure. Pas besoin d’un chat complexe pour le prototype."),
    ("06-recupere.png", "Écran 6 — C’est récupéré. L’annonce est clôturée, le tableau reste propre."),
]


def set_run(run, size=11, bold=False, italic=False, color=INK, font="Calibri"):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = font
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    rFonts.set(qn("w:eastAsia"), font)


def add_text(p, text, **kwargs):
    run = p.add_run(text)
    set_run(run, **kwargs)
    return run


def para(doc, text="", size=11, bold=False, italic=False, color=INK, space_after=8, space_before=0, font="Calibri"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    if text:
        add_text(p, text, size=size, bold=bold, italic=italic, color=color, font=font)
    return p


def heading(doc, text):
    p = para(doc, text, size=18, bold=True, color=BRICK, space_before=16, space_after=8, font="Georgia")
    return p


def mix(doc, parts, space_after=8):
    p = para(doc, space_after=space_after)
    for text, kwargs in parts:
        add_text(p, text, **kwargs)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2)
section.right_margin = Cm(2)

para(doc, "ESIG TECH ARENA 2026  ·  MINI-CHALLENGE DE PRÉSÉLECTION", size=10, bold=True, color=BRICK, space_after=4)
para(doc, "Récup’", size=36, bold=True, font="Georgia", space_after=6)
para(
    doc,
    "Le tableau du campus pour retrouver ce que tu as perdu, et emprunter ce qu’il te manque.",
    size=14,
    space_after=6,
)
para(doc, "EQUIPE_7  ·  BAKPASSIM Pouwedeo Light  ·  ABOLO-SEWOVI Ami Raphaëlla  ·  KOUYAKOUTOULI Godwin Marc", size=11, color=MUTED, space_after=2)
para(doc, "Prototype web cliquable  ·  17 septembre 2026", size=11, color=MUTED, space_after=14)

heading(doc, "A. Présentation du problème")
mix(doc, [
    ("Quel est le problème identifié ? ", {"bold": True, "size": 11}),
    ("Sur un campus, les objets bougent en continu : badges, chargeurs, calculatrices, écouteurs, clés USB, vestes. Quand un étudiant perd quelque chose, il n’a pas de lieu unique pour le signaler. Il écrit dans un groupe WhatsApp, demande autour de lui, passe parfois à l’accueil… et souvent il abandonne.", {"size": 11}),
])
mix(doc, [
    ("Celui qui trouve un objet est tout aussi bloqué : il ne sait pas à qui le rendre. L’objet finit dans un sac, un tiroir, ou à l’accueil sans que le propriétaire ne le sache. Le même chaos existe pour les prêts : avant un examen, on cherche une calculatrice dans la précipitation, sans savoir qui peut dépanner.", {"size": 11}),
])
mix(doc, [
    ("À qui ce problème se pose-t-il ? ", {"bold": True, "size": 11}),
    ("Aux étudiants au quotidien, surtout entre deux cours, à la BU, à la cafet, en amphi, et juste avant un partiel. L’accueil / la scolarité est concernée de façon secondaire : elle reçoit des objets sans canal simple pour les relier à leurs propriétaires.", {"size": 11}),
])
mix(doc, [
    ("Pourquoi ce problème mérite-t-il une solution ? ", {"bold": True, "size": 11}),
    ("Parce qu’il est fréquent, concret, et coûteux. Un badge à refaire, une calculatrice rachetée, des dizaines de minutes perdues à scroller un groupe, un stress inutile avant un examen. Les outils existants (WhatsApp, affichage papier, accueil) ne sont pas conçus pour ça : pas de photo structurée, pas de lieu, pas de statut « récupéré », pas de matching.", {"size": 11}),
])

heading(doc, "B. Présentation de la solution")
mix(doc, [
    ("Récup’", {"bold": True, "size": 11}),
    (" est un tableau numérique du campus. En 30 secondes, un étudiant peut signaler un objet trouvé, publier un objet perdu, ou proposer / demander un prêt.", {"size": 11}),
])
para(
    doc,
    "Chaque annonce a une catégorie, un lieu, une heure et un statut. Le propriétaire reconnaît son objet, contacte le trouveur, fixe un point de RDV sur le campus, puis clôture l’annonce. Le tableau reste propre. Le prêt suit la même logique : on voit ce qui est disponible maintenant, près de soi.",
)
para(doc, "Avantages", size=12, bold=True, space_after=4)
for item in [
    ("Rapide", "publier prend moins de temps qu’un message WhatsApp."),
    ("Lisible", "des cartes visuelles, pas un fil de 200 messages."),
    ("Utile deux fois", "objets trouvés et entraide (prêts)."),
    ("Fermé proprement", "« récupéré / rendu » évite les fausses pistes."),
    ("Ancré campus", "lieux réels, RDV entre étudiants."),
]:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    add_text(p, f"{item[0]} — ", bold=True, size=11)
    add_text(p, item[1], size=11)

heading(doc, "C. Fonctionnalités principales")
features = [
    ("J’ai trouvé", "catégorie, lieu, courte note. L’objet apparaît dans le fil Trouvé."),
    ("J’ai perdu", "description + lieu approximatif. La recherche est visible, plus noyée dans un groupe."),
    ("Je prête / j’emprunte", "« Calculatrice dispo jusqu’à 18h, Amphi B »."),
    ("C’est à moi", "proposition de RDV sur le campus (BU, cafet, hall)."),
    ("C’est récupéré / rendu", "l’annonce est clôturée. Le feed reste fiable."),
]
for title, body in features:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    add_text(p, f"{title} — ", bold=True, size=11)
    add_text(p, body, size=11)

heading(doc, "D. Public cible")
mix(doc, [
    ("Primaire : ", {"bold": True, "size": 11}),
    ("étudiants du campus, tous niveaux, au quotidien.", {"size": 11}),
])
mix(doc, [
    ("Secondaire : ", {"bold": True, "size": 11}),
    ("délégués, accueil, scolarité, qui peuvent y déposer les objets trouvés « officiels ».", {"size": 11}),
])
mix(doc, [
    ("Persona : ", {"bold": True, "size": 11}),
    ("Lina, 20 ans, 2e année. ", {"italic": True, "size": 11}),
    ("Elle a perdu son badge entre l’amphi et la cafet. Elle ouvre Récup’, voit une photo postée il y a 12 minutes par Karim, clique « C’est à moi », RDV à la BU.", {"size": 11}),
])

heading(doc, "E. Présentation du prototype")
para(
    doc,
    "Le prototype est une mini-application web cliquable. Les captures ci-dessous suivent le parcours de Lina et Karim.",
)
for name, caption in SHOTS:
    path = CAPS / name
    if path.exists():
        doc.add_picture(str(path), width=Cm(8.2))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
        last.paragraph_format.space_after = Pt(4)
    cap = para(doc, caption, size=10, italic=True, color=MUTED, space_after=12)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

heading(doc, "F. Liens complémentaires")
links = [
    "Prototype (fichier local) : prototype/index.html",
    "Présentation PowerPoint : livrables/Recup-presentation.pptx",
    "Présentation HTML (flèches clavier) : livrables/presentation.html",
    "Lien de démonstration en ligne : à coller ici si vous hébergez le proto (GitHub Pages, Netlify…)",
    "Dépôt GitHub : à coller ici si disponible",
]
for item in links:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    add_text(p, item, size=11)

para(
    doc,
    "Les liens sont un complément. Les captures ci-dessus sont le livrable visuel obligatoire.",
    size=11,
    italic=True,
    color=MUTED,
    space_before=8,
)

doc.save(OUT)
print(OUT)
